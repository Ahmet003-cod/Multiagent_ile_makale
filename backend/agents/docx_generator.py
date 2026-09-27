"""
Multi-Agent Workflow — Word Document Generator (15+ Sayfalık Akademik Monografi Formatı)

Markdown icerigini profesyonel, yayinevi / ders kitabi standartlarinda Word (.docx) dosyasina donusturur:
- Internetten toplanan gercek gorselleri ve diyagramlari indirir ve dokumana yerlestirir
- Gorsel alti aciklamalar ve kaynak atiflari (Sekil 1: ... Kaynak: ...)
- Renkli Bilgi Kutulari: [BUNU DA BIL?], [UYARI], [KIMDIR?], [NOT ALINIZ]
- Karsilastirmali Tablolar (Baslikli, renkli hucreli, cizgili)
- Sayfa ust/alt bilgileri ve kapak tasarimi
- Eksiksiz kaynakca ve bibliyografya
"""

import os
import re
import io
import html
import requests
from PIL import Image
from urllib.parse import urlparse
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


# Web tarayici basligi
BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
}
WIKI_HEADERS = {
    "User-Agent": "MultiAgentAcademicScholarBot/2.0 (scholar@multiagent.local)"
}


def add_hyperlink(paragraph, url: str, text: str, color_rgb: str = "0284C7", underline: bool = True):
    """
    Word belgesinde gercek, tiklanabilir ve tarayicida acilan OpenXML koprusu (hyperlink) olusturur.
    Kullanici Word'de tikladiginda veya Ctrl+Tik yaptiginda baglanti dogrudan tarayicida acilir.
    """
    try:
        clean_url = url.strip()
        part = paragraph.part
        r_id = part.relate_to(
            clean_url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True
        )

        escaped_text = html.escape(text.strip() if text else clean_url)
        u_val = "single" if underline else "none"
        hyperlink_xml = (
            f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<w:r {nsdecls("w")}>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
            f'<w:color w:val="{color_rgb}"/>'
            f'<w:u w:val="{u_val}"/>'
            f'<w:sz w:val="19"/>'  # 9.5 pt
            f'</w:rPr>'
            f'<w:t xml:space="preserve">{escaped_text}</w:t>'
            f'</w:r>'
            f'</w:hyperlink>'
        )
        hyperlink = parse_xml(hyperlink_xml)
        paragraph._p.append(hyperlink)
        return hyperlink
    except Exception:
        # Fallback to standard run if XML insertion encounters an edge case
        r = paragraph.add_run(text if text else url)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(2, 132, 199)
        r.font.underline = True
        return r


def extract_title_from_content(content: str, default_title: str) -> str:
    """Markdown icerisindeki gercek # Baslik metnini ceker."""
    if not content:
        return default_title
    for line in content.split("\n"):
        line = line.strip()
        if line.startswith("# ") and not line.startswith("## "):
            clean = line[2:].strip().strip("[]*#")
            clean = clean.split(" / ")[0].strip().strip("[]*")
            if len(clean) > 5 and not clean.lower().startswith("hedef"):
                return clean
    return default_title


def create_word_document(
    title: str, 
    content: str, 
    sources: list = None, 
    images: list = None,
    output_path: str = "output.docx",
    target_platform: str = "DergiPark"
) -> str:
    """Markdown icerigini gorselli, kutucuklu, tablolu Word belgesine donusturur."""
    doc = Document()

    # Sayfa Boyutu ve Kenar Bosluklari
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Gercek akademik basligi icerikten veya parametreden cek
    academic_title = extract_title_from_content(content, default_title=title)

    # Header / Footer
    _setup_header_footer(section, academic_title, target_platform)

    # Tipografi ve Stiller
    _setup_styles(doc)

    # Kapak Sayfasi
    _add_cover_page(doc, academic_title, target_platform)

    # Icerigi Parcala ve Ekle (Gorseller, Tablolar, Kutular)
    _parse_and_add_content(doc, content, images=images)

    # Kaynaklar Bolumu (Icerikte eksikse veya azsa dogrulanmis havuzu ekle)
    has_sources_in_content = any(k in content.lower() for k in ["# kaynaklar", "## kaynaklar", "### kaynaklar", "kaynakça", "bibliyografya", "references"])
    if sources and not has_sources_in_content:
        _add_sources_section(doc, sources)

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    doc.save(output_path)
    return output_path


def _setup_header_footer(section, title, target_platform="DergiPark"):
    """Ust ve alt bilgi ayarla."""
    is_blog = any(b in (target_platform or "").lower() for b in ["medium", "dev.to", "hashnode", "linkedin"])
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    label = "TEKNİK İNCELEME & BLOG" if is_blog else "AKADEMİK ARAŞTIRMA VE MONOGRAFİ"
    hrun = header_p.add_run(f"{label} | {title[:45]}")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(100, 116, 139)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_text = (
        f"Multi-Agent Workflow Sistemi ile Hazırlanmıştır — {target_platform} Formatı"
        if is_blog else
        "Multi-Agent Workflow Sistemi ile Derinlemesine Hazırlanmıştır — Akademik Hakemli Standart"
    )
    frun = footer_p.add_run(footer_text)
    frun.font.name = "Calibri"
    frun.font.size = Pt(8)
    frun.font.color.rgb = RGBColor(148, 163, 184)


def _setup_styles(doc):
    """Belge tipografisini ayarla."""
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.space_after = Pt(5)
    normal_style.paragraph_format.line_spacing = 1.25

    # Baslik 1
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(15, 23, 42)  # Cok Koyu Lacivert
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)

    # Baslik 2
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(2, 132, 199)  # Parlak Mavi
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    # Baslik 3
    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(13, 148, 136)  # Teal
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)


def _add_cover_page(doc, title, target_platform="DergiPark"):
    """Platform turune gore profesyonel kapak sayfasi."""
    is_blog = any(b in (target_platform or "").lower() for b in ["medium", "dev.to", "hashnode", "linkedin"])
    for _ in range(3):
        doc.add_paragraph("")

    badge_p = doc.add_paragraph()
    badge_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge_text = (
        f"SEKTÖREL TEKNİK İNCELEME & REHBER ({target_platform.upper()})"
        if is_blog else
        "AKADEMİK ARAŞTIRMA MONOGRAFİSİ & HAKEMLİ DERGİ MAKALESİ"
    )
    badge_r = badge_p.add_run(badge_text)
    badge_r.font.name = "Calibri"
    badge_r.font.size = Pt(12)
    badge_r.font.bold = True
    badge_r.font.color.rgb = RGBColor(2, 132, 199)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_r = title_p.add_run(title)
    title_r.font.name = "Calibri"
    title_r.font.size = Pt(24)
    title_r.font.bold = True
    title_r.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_before = Pt(14)
    title_p.paragraph_format.space_after = Pt(18)

    line_p = doc.add_paragraph()
    line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line_r = line_p.add_run("―" * 38)
    line_r.font.color.rgb = RGBColor(56, 189, 248)
    line_r.font.size = Pt(14)

    for _ in range(3):
        doc.add_paragraph("")

    info_tbl = doc.add_table(rows=1, cols=3)
    info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2, c3 = info_tbl.cell(0, 0), info_tbl.cell(0, 1), info_tbl.cell(0, 2)
    
    if is_blog:
        _set_cell_text(c1, "YAYIN FORMU", f"{target_platform}\nStorytelling & Kod Blokları", RGBColor(2, 132, 199))
        _set_cell_text(c2, "HEDEF KİTLE", "Teknoloji Profesyonelleri\nve Geliştiriciler", RGBColor(16, 185, 129))
        _set_cell_text(c3, "KAYNAKLAR", "Doğrulanmış Teknik Kaynaklar\nve Canlı Bağlantılar", RGBColor(124, 58, 237))
    else:
        _set_cell_text(c1, "YÖNTEM", "Çok Ajanlı İş Akışı\n(Researcher • Writer • Reviewer)", RGBColor(2, 132, 199))
        _set_cell_text(c2, "İÇERİK STANDARDI", "IMRaD Bilimsel Şablonu\nTablolar & Görseller & Kutular", RGBColor(16, 185, 129))
        _set_cell_text(c3, "KAYNAKÇA", "Doğrulanmış DOI / DergiPark\nMetin İçi Sıralı Atıf Sistemi", RGBColor(124, 58, 237))

    doc.add_page_break()


def _set_cell_text(cell, title, desc, color):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(title + "\n")
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = color
    r2 = p.add_run(desc)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(100, 116, 139)


def _insert_web_image(doc, img_url, alt_text, caption_source, figure_num, fallback_images=None, used_urls=None):
    """Web'den gorseli indirip PIL ile donusturerek Word dokumanina ekler. Her sekil numarasi icin FARKLI gorsel secer."""
    domain = urlparse(img_url).netloc
    img_added = False
    if used_urls is None:
        used_urls = set()

    # Adaylari akillica sirala: Oncelik kendi URL'sinde (eger daha once baska bir sekilde kullanilmadiysa)
    candidates = []
    if img_url and img_url not in used_urls:
        candidates.append(img_url)

    # Eger fallback_images varsa, once bu sekil numarasina karsilik gelen indeksteki gorseli aday yap
    if fallback_images:
        idx = max(0, figure_num - 1)
        if idx < len(fallback_images):
            target_fb = fallback_images[idx].get("url", "")
            if target_fb and target_fb not in candidates and target_fb not in used_urls:
                candidates.append(target_fb)
                # Baslik fallback'ten daha anlamli olabilir
                if not alt_text or "teknik" in alt_text.lower():
                    alt_text = fallback_images[idx].get("title", alt_text)
        
        # Diger kullanilmamis fallback gorsellerini ekle
        for fb in fallback_images:
            fb_url = fb.get("url", "")
            if fb_url and fb_url not in candidates and fb_url not in used_urls:
                candidates.append(fb_url)

    # Eger hic aday kalmadiysa, used_urls filtrelemesini kaldirarak siradaki gorseli al
    if not candidates and fallback_images:
        idx = (figure_num - 1) % len(fallback_images)
        candidates.append(fallback_images[idx].get("url", ""))

    for target_url in candidates:
        try:
            req_headers = WIKI_HEADERS if "wikimedia.org" in target_url.lower() or "wikipedia.org" in target_url.lower() else BROWSER_HEADERS
            resp = requests.get(target_url, headers=req_headers, timeout=10)
            if resp.status_code == 200 and len(resp.content) > 800:
                pil_img = Image.open(io.BytesIO(resp.content))
                
                # RGB JPEG formatina donustur
                buf = io.BytesIO()
                pil_img.convert("RGB").save(buf, format="JPEG", quality=90)
                buf.seek(0)

                # Gorseli ekle ve ortala
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(8)
                img_p.paragraph_format.space_after = Pt(3)
                
                run = img_p.add_run()
                run.add_picture(buf, width=Inches(5.4))
                img_added = True
                domain = urlparse(target_url).netloc
                used_urls.add(target_url)
                break
        except Exception:
            continue

    if img_added:
        # Gorsel Alti Aciklama ve Kaynak
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_before = Pt(2)
        cap_p.paragraph_format.space_after = Pt(10)

        r_num = cap_p.add_run(f"Şekil {figure_num}: ")
        r_num.bold = True
        r_num.font.size = Pt(9.5)
        r_num.font.color.rgb = RGBColor(15, 23, 42)

        r_title = cap_p.add_run(f"{alt_text or 'Teknik Şema / Mimari Diyagram'} ")
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(51, 65, 85)

        source_text = caption_source or f"(Görsel Kaynağı: {domain or 'Açık Bilim Arşivi'})"
        r_src = cap_p.add_run(f"— {source_text}")
        r_src.italic = True
        r_src.font.size = Pt(8.5)
        r_src.font.color.rgb = RGBColor(2, 132, 199)

    return img_added


def _parse_and_add_content(doc, content, images=None):
    """Markdown metnini parse edip gorseller, tablolar ve kutucuklar ile Word'e aktarir."""
    lines = content.split("\n")
    i = 0
    figure_counter = 1
    in_references_section = False
    used_image_urls = set()
    available_images = list(images) if images else []
    total_images_inserted = 0

    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # 1. Gorsel Yakalama: ![Alt Text](Image_URL)
        img_match = re.search(r"!\[(.*?)\]\((https?://[^\)\s]+)\)", line)
        if img_match:
            alt_text = img_match.group(1).strip()
            img_url = img_match.group(2).strip()
            
            # Sonraki satirda kaynak varsa oku
            caption_source = ""
            if i + 1 < len(lines) and ("Kaynak:" in lines[i + 1] or "Şekil" in lines[i + 1] or lines[i + 1].strip().startswith("*")):
                caption_source = lines[i + 1].strip().strip("*_")
                i += 1

            added = _insert_web_image(doc, img_url, alt_text, caption_source, figure_counter, fallback_images=available_images, used_urls=used_image_urls)
            if added:
                figure_counter += 1
                total_images_inserted += 1
            i += 1
            continue

        # 1.B: Eger yazar ![] yazmadi ama yalnizca 'Şekil X: ...' satiri yazdiysa
        fig_line_match = re.match(r"^[\*_]*Şekil\s*(\d+)[\:\.\-]\s*(.*?)[\*_]*$", line, re.IGNORECASE)
        if fig_line_match and available_images:
            fig_idx = figure_counter - 1
            if fig_idx < len(available_images):
                img_data = available_images[fig_idx]
                img_url = img_data.get("url", "")
                alt_text = img_data.get("title", "") or fig_line_match.group(2).strip()
                caption_source = img_data.get("source", "")
                
                added = _insert_web_image(doc, img_url, alt_text, caption_source, figure_counter, fallback_images=available_images, used_urls=used_image_urls)
                if added:
                    figure_counter += 1
                    total_images_inserted += 1
                i += 1
                continue

        # 2. Markdown Tablosu
        if line.startswith("|") and i + 1 < len(lines) and "---" in lines[i + 1]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _render_markdown_table(doc, table_lines)
            continue

        # 3. Ozel Bilgi Kutulari
        if line.startswith(">") or line.startswith("💡") or line.startswith("⚠️") or line.startswith("👤") or line.startswith("📌"):
            box_lines = []
            while i < len(lines) and (lines[i].strip().startswith(">") or lines[i].strip().startswith("💡") or lines[i].strip().startswith("⚠️") or lines[i].strip().startswith("👤") or lines[i].strip().startswith("📌") or (box_lines and lines[i].strip() and not lines[i].strip().startswith("#"))):
                curr = lines[i].strip().lstrip("> ").strip()
                if curr:
                    box_lines.append(curr)
                i += 1
                if i < len(lines) and not lines[i].strip():
                    break
            _render_callout_box(doc, box_lines)
            continue

        # 4. Basliklar
        if line.lower().startswith("#") and any(k in line.lower() for k in ["kaynakça", "kaynaklar", "references", "bibliyografya"]):
            in_references_section = True
            doc.add_page_break()
            doc.add_heading("KAYNAKÇA VE BİBLİYOGRAFYA", level=1)
        elif line.startswith("#### "):
            doc.add_heading(_clean_markdown(line[5:]), level=4)
        elif line.startswith("### "):
            doc.add_heading(_clean_markdown(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_clean_markdown(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_clean_markdown(line[2:]), level=1)

        # 5. Kaynakca Bolumundeki Tum Numaralandirilmis / Linkli Maddeler
        elif in_references_section and (
            re.match(r"^(\d+[\.\)]\s*)?\[\d+\]", line) or 
            re.match(r"^\d+[\.\)]\s+", line) or 
            re.match(r"^[\*\-•]\s*\[\d+\]", line) or
            "doi.org" in line.lower() or 
            "dergipark.org.tr" in line.lower() or
            "arxiv.org" in line.lower()
        ):
            # Eger numara '1. Yazar' seklindeyse '[1] Yazar' haline getir
            norm_line = line
            num_prefix_m = re.match(r"^(\d+)[\.\)]\s+(.*)", line)
            bracket_m = re.match(r"^(\d+[\.\)]\s*)?\[(\d+)\]\s*(.*)", line)
            bullet_bracket_m = re.match(r"^[\*\-•]\s*\[(\d+)\]\s*(.*)", line)
            
            if bracket_m:
                norm_line = f"[{bracket_m.group(2)}] {bracket_m.group(3)}"
            elif bullet_bracket_m:
                norm_line = f"[{bullet_bracket_m.group(1)}] {bullet_bracket_m.group(2)}"
            elif num_prefix_m:
                norm_line = f"[{num_prefix_m.group(1)}] {num_prefix_m.group(2)}"
                
            _render_reference_entry(doc, norm_line)

        # 6. Normal Metin Icerisindeki Atif / Kaynakca Maddesi: [1], [2], [14]
        elif re.match(r"^\[\d+\]\s*", line):
            _render_reference_entry(doc, line)

        # 7. Madde Isaretli Liste
        elif line.startswith("- ") or line.startswith("* ") or line.startswith("• "):
            text = _clean_markdown(line.lstrip("-*• "))
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_text(p, text)

        # 8. Numarali Liste
        elif re.match(r"^\d+[\.\)]\s", line):
            text = _clean_markdown(re.sub(r"^\d+[\.\)]\s*", "", line))
            p = doc.add_paragraph(style="List Number")
            _add_formatted_text(p, text)

        # 9. Normal Paragraf
        else:
            text = line
            if text:
                p = doc.add_paragraph()
                _add_formatted_text(p, text)

        i += 1

    # GARANTİ MEKANİZMASI: Eger metin bitti ve hic gorsel eklenmediyse,
    # eldeki available_images havuzundan gorselleri dokumana ekle!
    if total_images_inserted == 0 and available_images:
        for idx, img_data in enumerate(available_images[:4], 1):
            _insert_web_image(
                doc, 
                img_data.get("url", ""), 
                img_data.get("title", f"Teknik Şema {idx}"), 
                img_data.get("source", "Wikimedia Commons"), 
                idx, 
                fallback_images=available_images,
                used_urls=used_image_urls
            )


def _extract_clean_url(text: str) -> str:
    """Metin icerisindeki bozulmus, ic ice gecmis markdown linklerinden saf ve calisan URL'yi ceker."""
    # 1. Standart markdown linki ara: [text](http...)
    md_matches = re.findall(r'\[(.*?)\]\((https?://[^\)\s]+)\)', text)
    if md_matches:
        for display_txt, target_url in md_matches:
            target_clean = target_url.strip().rstrip(".,;:-/\\)]")
            second_http = target_clean.find("http", 4)
            if second_http != -1:
                target_clean = target_clean[:second_http]
            target_clean = re.sub(r'[\(\)\[\]\\\%]+.*$', '', target_clean).strip(".,;:-/\\")
            if "arxiv.org" in target_clean or "doi.org" in target_clean or "dergipark" in target_clean or len(target_clean) > 12:
                return target_clean

    # 2. Dogrudan URL ara
    raw_urls = re.findall(r'https?://[^\s\)\"\'\]\>]+', text)
    for u in raw_urls:
        u_clean = u.strip()
        second_http = u_clean.find("http", 4)
        if second_http != -1:
            u_clean = u_clean[:second_http]
        u_clean = re.sub(r'[\(\)\[\]\\\%]+.*$', '', u_clean).strip(".,;:-/\\")
        if "arxiv.org" in u_clean or "doi.org" in u_clean or "dergipark" in u_clean or len(u_clean) > 12:
            return u_clean

    # 3. DOI on eki tespiti: 10.xxxx/...
    doi_raw = re.search(r'10\.\d{4,9}/[-._;()/:A-Za-z0-9]+', text)
    if doi_raw:
        clean_doi = doi_raw.group(0).rstrip(".,;:-/\\")
        return f"https://doi.org/{clean_doi}"

    return ""


def _render_reference_entry(doc, line):
    """Akademik kaynakca maddesini asili girinti (hanging indent), [X] bold ve gercek tiklanabilir OpenXML DOI linkiyle ekler."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    match = re.match(r"^(\[\d+\])\s*(.*)", line)
    if not match:
        _add_formatted_text(p, _clean_markdown(line))
        return

    num_tag = match.group(1)
    body = match.group(2).strip()

    # Numara: [1]
    num_run = p.add_run(f"{num_tag} ")
    num_run.bold = True
    num_run.font.name = "Calibri"
    num_run.font.size = Pt(10)
    num_run.font.color.rgb = RGBColor(15, 23, 42)

    # Saf, calisan ve dogrulanmis URL cikar
    clean_url = _extract_clean_url(body)
    if clean_url:
        if clean_url.startswith("http://"):
            clean_url = "https://" + clean_url[7:]

        # Body metnindeki bozuk [. Erisim: [http...](http...) kisimlarini ayikla
        text_before_doi = re.sub(r'\[?\.\s*(?:Erişim|DOI|URL)\s*:\s*.*', '', body, flags=re.IGNORECASE).strip()
        text_before_doi = re.sub(r'\[https?://.*', '', text_before_doi).strip()
        text_before_doi = text_before_doi.rstrip(".,;:- ")

        _add_formatted_text(p, _clean_markdown(text_before_doi))
        p.add_run(". ")

        doi_label = "DOI: " if "doi.org" in clean_url else "Erişim: "
        doi_tag = p.add_run(doi_label)
        doi_tag.font.bold = True
        doi_tag.font.size = Pt(9.5)
        doi_tag.font.color.rgb = RGBColor(71, 85, 105)

        # GERÇEK AÇILABİLİR OPENXML WORD KÖPRÜSÜ (Clickable Hyperlink)
        add_hyperlink(p, clean_url, clean_url, color_rgb="0284C7", underline=True)
    else:
        _add_formatted_text(p, _clean_markdown(body))


def _render_callout_box(doc, box_lines):
    """Referans PDF formatinda renkli, kenarlikli bilgi kutusu."""
    full_text = "\n".join(box_lines)
    upper = full_text.upper()

    if "TARİHSEL" in upper or "TARIHSEL" in upper or "KİLOMETRE TAŞI" in upper or "KILOMETRE TASI" in upper:
        bg_hex = "FEF3C7"      # Sicak sari / altin
        border_hex = "B45309"  # Koyu Kehribar
        title_color = RGBColor(180, 83, 9)
        icon = "📜 TARİHSEL KİLOMETRE TAŞI:"
    elif "PRO TIP" in upper or "İPUCU" in upper or "IPUCU" in upper:
        bg_hex = "ECFDF5"      # Acik Yesil
        border_hex = "059669"  # Zumrut Yesili
        title_color = RGBColor(5, 150, 105)
        icon = "💡 PRO TIP / UZMAN TAVSİYESİ:"
    elif "UYARI" in upper or "DİKKAT" in upper:
        bg_hex = "FFFBEB"      # Acik Amber
        border_hex = "D97706"  # Amber
        title_color = RGBColor(180, 83, 9)
        icon = "⚠️ UYARI / DİKKAT:"
    elif "KİMDİR" in upper or "KIMDIR" in upper:
        bg_hex = "F5F3FF"      # Acik Mor
        border_hex = "7C3AED"  # Mor
        title_color = RGBColor(109, 40, 217)
        icon = "👤 KİMDİR / ÖNCÜ TEORİSYEN?"
    elif "BUNU DA BİL" in upper or "BUNUDA BIL" in upper or "DERİN ANALİZ" in upper or "DERIN ANALIZ" in upper:
        bg_hex = "E0F2FE"      # Acik Mavi
        border_hex = "0284C7"  # Mavi
        title_color = RGBColor(3, 105, 161)
        icon = "💡 DERİN TEKNİK ANALİZ:"
    else:
        bg_hex = "F8FAFC"      # Acik Gri
        border_hex = "475569"  # Slate
        title_color = RGBColor(51, 65, 85)
        icon = "📌 NOT ALINIZ:"

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.2)

    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()

    # Arka plan rengi
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    tcPr.append(shd)

    # Sol kenarlik
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    # Baslik
    p_head = cell.paragraphs[0]
    p_head.paragraph_format.space_after = Pt(3)
    r_head = p_head.add_run(icon)
    r_head.bold = True
    r_head.font.size = Pt(10.5)
    r_head.font.color.rgb = title_color

    for line in box_lines:
        clean = _clean_markdown(line)
        if any(h in clean.upper() for h in ["UYARI", "KİMDİR", "KIMDIR", "BUNU DA BİL", "NOT ALINIZ"]) and len(clean) < 25:
            continue
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _add_formatted_text(p, clean)

    doc.add_paragraph("")


def _render_markdown_table(doc, table_lines):
    """Markdown tablolarini renkli ve cizgili Word tablosuna donusturur."""
    rows_data = []
    for line in table_lines:
        if "---" in line:
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells:
            rows_data.append(cells)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            if c_idx < num_cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = _clean_markdown(val)
                tcPr = cell._tc.get_or_add_tcPr()

                if r_idx == 0:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
                    tcPr.append(shd)
                    p = cell.paragraphs[0]
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)
                else:
                    if r_idx % 2 == 1:
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                        tcPr.append(shd)
                    p = cell.paragraphs[0]
                    for r in p.runs:
                        r.font.size = Pt(9.5)

    doc.add_paragraph("")


def _add_formatted_text(paragraph, text):
    """Metin ici **bold**, *italic* ve [link metni](url) formatlamasini cozer."""
    # Once markdown baglantilarini ayir: [Metin](URL)
    link_parts = re.split(r"(\[[^\]]+\]\(https?://[^\)]+\))", text)
    for lpart in link_parts:
        m = re.match(r"^\[([^\]]+)\]\((https?://[^\)]+)\)$", lpart)
        if m:
            label = m.group(1)
            url = m.group(2)
            add_hyperlink(paragraph, url, label, color_rgb="0284C7", underline=True)
        else:
            parts = re.split(r"(\*\*.*?\*\*)", lpart)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = paragraph.add_run(part[2:-2])
                    run.bold = True
                else:
                    italic_parts = re.split(r"(\*.*?\*)", part)
                    for ipart in italic_parts:
                        if ipart.startswith("*") and ipart.endswith("*") and len(ipart) > 2:
                            run = paragraph.add_run(ipart[1:-1])
                            run.italic = True
                        else:
                            paragraph.add_run(ipart)


def _clean_markdown(text):
    return text.strip()


def _add_sources_section(doc, sources):
    """Dogrulanmis kaynaklar icin asili girintili ve tiklanabilir kaynakca bolumu."""
    doc.add_page_break()
    doc.add_heading("KAYNAKÇA VE BİBLİYOGRAFYA", level=1)

    for i, source in enumerate(sources, 1):
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.8)
        para.paragraph_format.first_line_indent = Cm(-0.8)
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.15

        run_num = para.add_run(f"[{i}] ")
        run_num.bold = True
        run_num.font.name = "Calibri"
        run_num.font.size = Pt(10)
        
        title = source.get("title", f"Akademik Kaynak {i}")
        run_title = para.add_run(f"{title}. ")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10)
        
        url = source.get("url", "")
        if url:
            lbl = "DOI: " if "doi.org" in url else "Erişim: "
            run_lbl = para.add_run(lbl)
            run_lbl.bold = True
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.color.rgb = RGBColor(71, 85, 105)
            add_hyperlink(para, url, url, color_rgb="0284C7", underline=True)


def extract_sources_from_text(text):
    sources = []
    urls = re.findall(r"https?://[^\s\)\"'>]+", text)
    for url in urls:
        domain = urlparse(url).netloc
        title = domain or url
        if url not in [s.get("url") for s in sources]:
            sources.append({"title": title, "url": url})
    return sources
